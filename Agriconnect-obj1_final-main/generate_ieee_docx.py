from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


INPUT = Path("IEEE_PAPER_DRAFT.md")
OUTPUT = Path("IEEE_PAPER_DRAFT.docx")

AUTHORS = [
    {
        "name": "Yashaswini Manyam",
        "dept": "Department of Computer Science and Engineering",
        "school": "Amrita School of Computing, Bangalore",
        "university": "Amrita Vishwa Vidyapeetham, India",
        "email": "bl.en.u4aie22036@bl.students.amrita.edu",
    },
    {
        "name": "Murari Nallamalli",
        "dept": "Department of Computer Science and Engineering",
        "school": "Amrita School of Computing, Bangalore",
        "university": "Amrita Vishwa Vidyapeetham, India",
        "email": "bl.en.u4aie22042@bl.students.amrita.edu",
    },
    {
        "name": "Hemanth Saga",
        "dept": "Department of Computer Science and Engineering",
        "school": "Amrita School of Computing, Bangalore",
        "university": "Amrita Vishwa Vidyapeetham, India",
        "email": "bl.en.u4aie22049@bl.students.amrita.edu",
    },
    {
        "name": "Thanuj Raja",
        "dept": "Department of Computer Science and Engineering",
        "school": "Amrita School of Computing, Bangalore",
        "university": "Amrita Vishwa Vidyapeetham, India",
        "email": "bl.en.u4aie22140@bl.students.amrita.edu",
    },
    {
        "name": "Radha D.",
        "dept": "Department of Computer Science and Engineering",
        "school": "Amrita School of Computing, Bangalore",
        "university": "Amrita Vishwa Vidyapeetham, India",
        "email": "d_radha@blr.amrita.edu",
    },
]


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(8)


def set_doc_defaults(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.0

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(0.625)
        section.right_margin = Inches(0.625)


def set_columns(section, count=2, space_twips=360):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(count))
    cols.set(qn("w:space"), str(space_twips))


def clear_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "nil")
        borders.append(tag)
    tbl_pr.append(borders)


def set_table_width(table, width_pct=100):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(width_pct * 50))
    tbl_w.set(qn("w:type"), "pct")
    tbl_pr.append(tbl_w)


def add_paragraph(doc, text, align=None, bold=False, italic=False, size=10):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0

    # Very small Markdown inline handling for bold labels.
    if text.startswith("**") and ":**" in text:
        label, rest = text.split(":**", 1)
        r = p.add_run(label.replace("**", "") + ":")
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        r2 = p.add_run(rest.replace("**", ""))
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(size)
    else:
        run = p.add_run(text.replace("**", ""))
        run.bold = bold
        run.italic = italic
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
    return p


def add_author_cell(cell, author):
    cell.text = ""
    lines = [
        (author["name"], False),
        (author["dept"], True),
        (author["school"], True),
        (author["university"], False),
        (author["email"], False),
    ]
    for idx, (text, italic) in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(7.5)
        run.italic = italic


def add_author_block(doc):
    top = doc.add_table(rows=1, cols=3)
    clear_table_borders(top)
    set_table_width(top)
    for idx, author in enumerate(AUTHORS[:3]):
        add_author_cell(top.cell(0, idx), author)
    doc.add_paragraph()

    bottom = doc.add_table(rows=1, cols=2)
    clear_table_borders(bottom)
    set_table_width(bottom)
    for idx, author in enumerate(AUTHORS[3:]):
        add_author_cell(bottom.cell(0, idx), author)
    doc.add_paragraph()


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 2 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8 if level == 2 else 5)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10 if level == 2 else 9)
    return p


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            set_cell_text(table.cell(r_idx, c_idx), value.strip(), bold=(r_idx == 0))
    doc.add_paragraph()


def add_figure(doc, image_path, caption):
    if not image_path.exists():
        add_paragraph(doc, caption + " [image missing]", WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=8)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(3.15))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run(caption)
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(8)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        parts = [p.strip() for p in lines[i].strip().strip("|").split("|")]
        if not all(set(p) <= {"-", ":"} for p in parts):
            rows.append(parts)
        i += 1
    return rows, i


def build_doc():
    lines = INPUT.read_text(encoding="utf-8").splitlines()
    doc = Document()
    set_doc_defaults(doc)

    title = lines[0].replace("# ", "").strip()
    add_paragraph(doc, title, WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=18)
    add_author_block(doc)

    i = 1
    # Skip Markdown author block; the DOCX uses the IEEE-style grid above.
    while i < len(lines):
        line = lines[i].strip()
        if line == "## Abstract":
            break
        i += 1

    # Start IEEE-like two-column body before abstract, matching the accepted sample.
    body_section = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(body_section, 2)

    add_heading(doc, "Abstract", 2)
    i += 1
    while i < len(lines):
        line = lines[i].strip()
        if line.startswith("**Keywords:**") or line.startswith("**Index Terms**"):
            add_paragraph(doc, line, None, size=9)
            i += 1
            break
        if line:
            add_paragraph(doc, line, None, size=9)
        i += 1

    in_code = False
    code_buffer = []
    pending_table_title = None

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()

        if not line:
            i += 1
            continue

        if line.startswith("```"):
            if in_code:
                for item in code_buffer:
                    add_paragraph(doc, item, None, size=8)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        if line.startswith("## References"):
            add_heading(doc, "References", 2)
            i += 1
            continue

        if line.startswith("## "):
            add_heading(doc, line.replace("## ", ""), 2)
            i += 1
            continue

        if line.startswith("### "):
            add_heading(doc, line.replace("### ", ""), 3)
            i += 1
            continue

        if line.startswith("**Table"):
            pending_table_title = line.replace("**", "")
            add_paragraph(doc, pending_table_title, WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=8)
            i += 1
            continue

        if line.startswith("![") and "](" in line and line.endswith(")"):
            caption = line[2 : line.index("](")]
            image_path = Path(line[line.index("](") + 2 : -1])
            add_figure(doc, image_path, caption)
            i += 1
            continue

        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                add_table(doc, rows)
            pending_table_title = None
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style=None)
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.first_line_indent = Inches(-0.1)
            r = p.add_run("- " + line[2:])
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            i += 1
            continue

        add_paragraph(doc, line, None, size=10)
        i += 1

    try:
        doc.save(OUTPUT)
    except PermissionError:
        alt_output = Path("IEEE_PAPER_DRAFT_ACCEPTED_STYLE.docx")
        try:
            doc.save(alt_output)
            print(f"{OUTPUT} is open or locked. Wrote {alt_output} instead.")
        except PermissionError:
            stamped = Path(f"IEEE_PAPER_DRAFT_ACCEPTED_STYLE_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
            doc.save(stamped)
            print(f"{OUTPUT} and {alt_output} are open or locked. Wrote {stamped} instead.")
        return


if __name__ == "__main__":
    build_doc()
    print(f"Wrote {OUTPUT}")


